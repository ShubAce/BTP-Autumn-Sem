"""
=============================================================================
  Wheel Wander Statistical Analysis  —  BTP IIT Kharagpur
=============================================================================
  Run AFTER pipeline.py has generated the CSV files.

  Usage:
    python analyze.py --csv results/*.csv --output-dir analysis_output

  Produces:
    • analysis_output/wheel_wander_report.xlsx   — full stats
    • analysis_output/figures/                   — PNG charts
=============================================================================
"""

from __future__ import annotations
import argparse, sys
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
from pathlib import Path
from typing import List

import numpy as np

try:
    import pandas as pd
    PANDAS_OK = True
except ImportError:
    PANDAS_OK = False; print("[ERROR] pip install pandas"); sys.exit(1)

try:
    import matplotlib.pyplot as plt
    import matplotlib.ticker as mticker
    from matplotlib.gridspec import GridSpec
    import seaborn as sns
    MPL_OK = True
except ImportError:
    MPL_OK = False; print("[WARN] pip install matplotlib seaborn")

try:
    from scipy import stats as sp_stats
    SCIPY_OK = True
except ImportError:
    SCIPY_OK = False

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import BarChart, Reference
    XLSX_OK = True
except ImportError:
    XLSX_OK = False

# ── Style ─────────────────────────────────────────────────────────────────
PALETTE = {
    "multi_axle_truck": "#D90429",
    "truck":            "#E63946",
    "bus":              "#F4A261",
    "pickup":           "#E76F51",
    "car":              "#457B9D",
    "autorickshaw":     "#E9C46A",
    "motorcycle":       "#2A9D8F",
    "bicycle":          "#9B5DE5",
    "rider":            "#A8DADC",
    "other":            "#B0B0B0",
}
BG      = "#0F1924"
FG      = "#E8EDF2"
ACCENT  = "#00C2D4"
C_WHEEL_L = (0, 200, 255)
C_WHEEL_R = (255, 80, 200)

plt.rcParams.update({
    "figure.facecolor":  BG,
    "axes.facecolor":    "#1A2535",
    "axes.edgecolor":    "#2E3E54",
    "axes.labelcolor":   FG,
    "xtick.color":       FG,
    "ytick.color":       FG,
    "text.color":        FG,
    "grid.color":        "#2E3E54",
    "grid.linewidth":    0.6,
    "legend.facecolor":  "#1A2535",
    "legend.edgecolor":  "#2E3E54",
    "font.family":       "DejaVu Sans",
})

# ═══════════════════════════════════════════════════════════════════════════
def load_csvs(csv_paths: List[str]) -> pd.DataFrame:
    dfs = []
    for p in csv_paths:
        try:
            dfs.append(pd.read_csv(p))
            print(f"[LOAD] {p}  ({len(dfs[-1])} rows)")
        except Exception as e:
            print(f"[SKIP] {p} — {e}")
    if not dfs:
        sys.exit("[ERROR] No valid CSV files found.")
    df = pd.concat(dfs, ignore_index=True)
    # Normalise column names
    df.columns = [c.strip() for c in df.columns]
    rename = {
        "LeftWheelOffset(m)":  "lw_off",
        "RightWheelOffset(m)": "rw_off",
        "CenterOffset(m)":     "ctr_off",
        "VehicleType":         "type",
        "VehicleID":           "vid_id",
        "Frame":               "frame",
        "Timestamp(s)":        "ts",
        "Confidence":          "conf",
        "Video":               "video",
    }
    df = df.rename(columns={k: v for k,v in rename.items() if k in df.columns})
    # cm versions
    for col in ("lw_off", "rw_off", "ctr_off"):
        if col in df.columns:
            df[f"{col}_cm"] = df[col] * 100.0
    return df

# ═══════════════════════════════════════════════════════════════════════════
def stats_table(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for vt in ["ALL"] + sorted(df["type"].dropna().unique().tolist()):
        sub = df if vt == "ALL" else df[df["type"] == vt]
        if sub.empty: continue
        for col, label in [
            ("lw_off_cm",  "Left-Wheel Offset (cm)"),
            ("rw_off_cm",  "Right-Wheel Offset (cm)"),
            ("ctr_off_cm", "Center Offset (cm)"),
        ]:
            if col not in sub.columns: continue
            s = sub[col].dropna()
            if s.empty: continue
            row = {
                "Vehicle Type": vt,
                "Metric":       label,
                "Count":        len(s),
                "Mean (cm)":    round(s.mean(), 3),
                "Std (cm)":     round(s.std(), 3),
                "Min (cm)":     round(s.min(), 3),
                "Max (cm)":     round(s.max(), 3),
                "Median (cm)":  round(s.median(), 3),
                "P5 (cm)":      round(s.quantile(0.05), 3),
                "P95 (cm)":     round(s.quantile(0.95), 3),
            }
            if SCIPY_OK and len(s) > 3:
                vals = pd.to_numeric(s, errors="coerce").dropna().astype(float).to_numpy()
                if vals.size > 3:
                    sigma = float(np.std(vals))
                    if sigma > 1e-12:
                        z = (vals - float(np.mean(vals))) / sigma
                        _, kp = sp_stats.kstest(z.astype(float), 'norm')
                        row["KS p-value"] = round(float(kp), 4)
                        row["Normality"]  = "Yes" if kp > 0.05 else "No"
                    else:
                        row["KS p-value"] = None
                        row["Normality"]  = "Flat"
            rows.append(row)
    return pd.DataFrame(rows)

# ═══════════════════════════════════════════════════════════════════════════
def make_figures(df: pd.DataFrame, out_dir: Path):
    if not MPL_OK: return
    fig_dir = out_dir / "figures"
    fig_dir.mkdir(parents=True, exist_ok=True)

    trucks = df[df["type"].isin(["truck", "bus", "heavy_truck", "multi_axle_truck", "pickup_lcv", "minibus", "pickup"])]

    # ── Fig 1: Distribution overview (3 panels) ────────────────────────
    fig, axes = plt.subplots(1, 3, figsize=(18, 6))
    fig.suptitle("Lateral Offset Distribution — All Vehicles",
                 fontsize=15, color=FG, y=1.01)
    for ax, (col, lbl) in zip(axes, [
        ("lw_off_cm",  "Left-Wheel Offset (cm)"),
        ("rw_off_cm",  "Right-Wheel Offset (cm)"),
        ("ctr_off_cm", "Center Offset (cm)"),
    ]):
        if col not in df.columns: continue
        for vt, grp in df.groupby("type"):
            color = PALETTE.get(vt, "#888")
            sns.kdeplot(grp[col].dropna(), ax=ax, label=vt,
                        color=color, linewidth=2, fill=True, alpha=0.15)
        ax.set_xlabel(lbl, fontsize=11)
        ax.set_ylabel("Density", fontsize=11)
        ax.legend(fontsize=9)
        ax.grid(True, alpha=0.3)
    plt.tight_layout()
    p = fig_dir / "01_distribution_overview.png"
    plt.savefig(p, dpi=150, bbox_inches="tight", facecolor=BG)
    plt.close(); print(f"[FIG] {p}")

    # ── Fig 2: Truck wheel-wander (left vs right vs center) ───────────
    if not trucks.empty and "ctr_off_cm" in trucks.columns:
        fig, axes = plt.subplots(1, 3, figsize=(18, 6))
        fig.suptitle("Trucks & Buses — Wheel Wander Analysis", fontsize=15, color=FG)
        for ax, (col, lbl, col_c) in zip(axes, [
            ("lw_off_cm",  "Left Wheel",   C_WHEEL_L[::-1]),
            ("rw_off_cm",  "Right Wheel",  C_WHEEL_R[::-1]),
            ("ctr_off_cm", "Center",       "#FFD700"),
        ]):
            if col not in trucks.columns: continue
            s = trucks[col].dropna()
            color_hex = "#{:02X}{:02X}{:02X}".format(*[int(x) for x in
                            [col_c[0] if isinstance(col_c, tuple) else 0,
                             0, 0]]) if False else "#00C2D4"
            ax.hist(s, bins=30, color=ACCENT, edgecolor="#0F1924",
                    alpha=0.85, density=True)
            if SCIPY_OK and len(s) > 5:
                mu, sigma = s.mean(), s.std()
                xs = np.linspace(s.min(), s.max(), 200)
                ax.plot(xs, sp_stats.norm.pdf(xs, mu, sigma),
                        "r--", lw=2, label=f"Normal μ={mu:.1f}")
            ax.axvline(s.mean(), color="yellow", lw=1.5,
                       linestyle="--", label=f"Mean={s.mean():.1f}")
            ax.set_xlabel(f"{lbl} Offset (cm)", fontsize=11)
            ax.set_ylabel("Density", fontsize=11)
            ax.legend(fontsize=9); ax.grid(True, alpha=0.3)
            ax.set_title(f"σ={s.std():.2f} cm  n={len(s)}", fontsize=10, color=FG)
        plt.tight_layout()
        p = fig_dir / "02_truck_wheel_wander.png"
        plt.savefig(p, dpi=150, bbox_inches="tight", facecolor=BG)
        plt.close(); print(f"[FIG] {p}")

    # ── Fig 3: Box-plots by vehicle type ──────────────────────────────
    if "ctr_off_cm" in df.columns:
        fig, ax = plt.subplots(figsize=(12, 6))
        order = sorted(df["type"].dropna().unique())
        colors = [PALETTE.get(v, "#888") for v in order]
        bplot  = ax.boxplot(
            [df[df["type"]==v]["ctr_off_cm"].dropna().values for v in order],
            patch_artist=True, notch=False,
            medianprops=dict(color="yellow", lw=2),
            whiskerprops=dict(color=FG), capprops=dict(color=FG),
            flierprops=dict(marker="o", color=ACCENT, alpha=0.4, markersize=4),
        )
        for patch, color in zip(bplot["boxes"], colors):
            patch.set_facecolor(color); patch.set_alpha(0.7)
        ax.set_xticklabels(order, fontsize=11)
        ax.set_ylabel("Center Offset (cm)", fontsize=12)
        ax.set_title("Center Offset by Vehicle Type", fontsize=14, color=FG)
        ax.grid(True, axis="y", alpha=0.3)
        ax.axhline(0, color="red", lw=1, linestyle="--", label="Right Boundary")
        ax.legend(fontsize=10)
        plt.tight_layout()
        p = fig_dir / "03_boxplot_by_type.png"
        plt.savefig(p, dpi=150, bbox_inches="tight", facecolor=BG)
        plt.close(); print(f"[FIG] {p}")

    # ── Fig 4: Lane utilisation heat-map (truck only) ─────────────────
    if (not trucks.empty and
            "lw_off_cm" in trucks.columns and "rw_off_cm" in trucks.columns):
        fig, ax = plt.subplots(figsize=(10, 7))
        ax.hist2d(trucks["lw_off_cm"].dropna(),
                  trucks["rw_off_cm"].dropna(),
                  bins=30, cmap="plasma")
        ax.set_xlabel("Left-Wheel Offset (cm)", fontsize=12)
        ax.set_ylabel("Right-Wheel Offset (cm)", fontsize=12)
        ax.set_title("Truck Lane Utilisation Heat-Map\n(Left vs Right Wheel)", fontsize=13, color=FG)
        cbar = plt.colorbar(ax.collections[0], ax=ax)
        cbar.set_label("Count", color=FG)
        plt.tight_layout()
        p = fig_dir / "04_lane_utilisation_heatmap.png"
        plt.savefig(p, dpi=150, bbox_inches="tight", facecolor=BG)
        plt.close(); print(f"[FIG] {p}")

    # ── Fig 5: Time-series of center offset ───────────────────────────
    if "ts" in df.columns and "ctr_off_cm" in df.columns:
        fig, ax = plt.subplots(figsize=(16, 5))
        for vt, grp in df.groupby("type"):
            color = PALETTE.get(vt, "#888")
            grp_s = grp.sort_values("ts")
            ax.scatter(grp_s["ts"], grp_s["ctr_off_cm"],
                       color=color, label=vt, s=20, alpha=0.7)
        ax.set_xlabel("Timestamp (s)", fontsize=12)
        ax.set_ylabel("Center Offset (cm)", fontsize=12)
        ax.set_title("Center Offset Time-Series", fontsize=14, color=FG)
        ax.axhline(0, color="red", lw=1, linestyle="--")
        ax.legend(fontsize=9); ax.grid(True, alpha=0.3)
        plt.tight_layout()
        p = fig_dir / "05_center_offset_timeseries.png"
        plt.savefig(p, dpi=150, bbox_inches="tight", facecolor=BG)
        plt.close(); print(f"[FIG] {p}")

    # ── Fig 6: Vehicle count pie ───────────────────────────────────────
    vc = df["type"].value_counts()
    if not vc.empty:
        fig, ax = plt.subplots(figsize=(7, 7))
        colors  = [PALETTE.get(v, "#888") for v in vc.index]
        pie_out = ax.pie(
            vc.values, labels=vc.index, colors=colors,
            autopct="%1.1f%%", startangle=140,
            textprops={"color": FG},
        )
        if hasattr(pie_out, "autotexts"):
            autotexts = pie_out.autotexts
        elif isinstance(pie_out, (tuple, list)) and len(pie_out) == 3:
            autotexts = pie_out[2]
        else:
            autotexts = []
        for at in autotexts:
            at.set_fontsize(10); at.set_color(BG)
        ax.set_title("Vehicle Type Distribution", fontsize=14, color=FG)
        plt.tight_layout()
        p = fig_dir / "06_vehicle_distribution.png"
        plt.savefig(p, dpi=150, bbox_inches="tight", facecolor=BG)
        plt.close(); print(f"[FIG] {p}")

# ═══════════════════════════════════════════════════════════════════════════
def export_excel(df: pd.DataFrame, stats: pd.DataFrame, out_path: Path):
    if not XLSX_OK:
        print("[WARN] openpyxl not available — skipping Excel export")
        return

    wb = openpyxl.Workbook()
    thin  = Side(style="thin",   color="CCCCCC")
    bdr   = Border(left=thin, right=thin, top=thin, bottom=thin)
    H_FILL = PatternFill("solid", start_color="1A2744")
    A_FILL = PatternFill("solid", start_color="EEF2FA")

    # ─ Sheet 1: Raw data ─────────────────────────────────────────────
    ws = wb.active; ws.title = "Raw Data"
    cols = [c for c in df.columns if c in [
        "video","type","vid_id","frame","ts",
        "lw_off","rw_off","ctr_off",
        "lw_off_cm","rw_off_cm","ctr_off_cm","conf"]]
    ws.append(cols)
    for c_i, col in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c_i)
        cell.font  = Font(bold=True, color="FFFFFF")
        cell.fill  = H_FILL; cell.border = bdr
        cell.alignment = Alignment(horizontal="center")
    for i, (_, row) in enumerate(df[cols].iterrows()):
        ws.append(row.tolist())
        fill = A_FILL if i % 2 == 0 else PatternFill("solid", start_color="FFFFFF")
        for c_i in range(1, len(cols)+1):
            cell = ws.cell(row=2+i, column=c_i)
            cell.fill = fill; cell.border = bdr
            cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"

    # ─ Sheet 2: Statistics ────────────────────────────────────────────
    ws2 = wb.create_sheet("Statistics")
    stat_cols = list(stats.columns)
    ws2.append(stat_cols)
    for c_i, col in enumerate(stat_cols, 1):
        cell = ws2.cell(row=1, column=c_i)
        cell.font  = Font(bold=True, color="FFFFFF")
        cell.fill  = H_FILL; cell.border = bdr
        cell.alignment = Alignment(horizontal="center")
    for i, (_, row) in enumerate(stats.iterrows()):
        ws2.append(row.tolist())
        fill = A_FILL if i % 2 == 0 else PatternFill("solid", start_color="FFFFFF")
        for c_i in range(1, len(stat_cols)+1):
            ws2.cell(row=2+i, column=c_i).fill   = fill
            ws2.cell(row=2+i, column=c_i).border = bdr
            ws2.cell(row=2+i, column=c_i).alignment = Alignment(horizontal="center")
    for col in ws2.columns:
        ws2.column_dimensions[col[0].column_letter].width = 18
    ws2.freeze_panes = "A2"

    wb.save(out_path)
    print(f"[XLSX] {out_path}")

# ═══════════════════════════════════════════════════════════════════════════
def main():
    ap = argparse.ArgumentParser(
        description="Wheel Wander Statistical Analysis — BTP IIT Kharagpur",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter
    )
    ap.add_argument("--csv",        nargs="+", required=True,
                    help="One or more CSV files output by pipeline.py")
    ap.add_argument("--output-dir", default="analysis_output",
                    help="Directory to write results into")
    args = ap.parse_args()

    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    df    = load_csvs(args.csv)
    stats = stats_table(df)

    print("\n" + "="*60)
    print("  Summary Statistics")
    print("="*60)
    print(stats.to_string(index=False))
    print("="*60 + "\n")

    stats.to_csv(out_dir / "statistics.csv", index=False)
    make_figures(df, out_dir)
    export_excel(df, stats, out_dir / "wheel_wander_report.xlsx")
    print(f"\n[DONE] All outputs in: {out_dir.resolve()}")

if __name__ == "__main__":
    main()
