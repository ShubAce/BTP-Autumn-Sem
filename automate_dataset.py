r"""
=============================================================================
  Automated Resumable Video-Processing System  —  BTP IIT Kharagpur
=============================================================================
  Processes videos in a strict folder hierarchy (Section -> Date -> Session -> Video)
  and records all results in a central, restart-safe Master Excel report.

  Usage:
    python automate_dataset.py --root "f:\BTP\NH-19 1st and 2nd May"
                               --output-dir outputs
                               --calib-file calib.json
                               --model yolo11x --conf 0.20 --imgsz 640 --skip-frames 2

  Features:
    • Dynamic Hierarchy Traversal: Section (LHS/RHS) -> Date -> Session -> Video
    • Single Master Excel File: outputs/master_dataset_report.xlsx
    • Fault-Tolerant Checkpoint: outputs/checkpoint.json
    • Seamless Resume: Automatically skips completed videos and retries interrupted ones
    • Integrated Analytics & Plots: Runs analyze.py across all processed videos
=============================================================================
"""

from __future__ import annotations
import argparse
import csv
import datetime
import json
import os
import re
import signal
import subprocess
import sys
import time
import traceback
from pathlib import Path
from typing import Dict, List, Optional, Tuple

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    XLSX_OK = True
except ImportError:
    XLSX_OK = False

VIDEO_EXTS = {".ts", ".mp4", ".avi", ".mkv", ".mov", ".mpeg", ".mpg"}

# The RHS footage from 2nd May covers three standard 3.5 m lanes.  Calibration
# dimensions are properties of a section/date group, not an individual video.
THREE_LANE_GROUPS = {("RHS", "2nd May")}


def natural_key(text: str) -> list:
    """Natural alphanumeric sort key (e.g., '1st May' before '2nd May', '0000-0100' before '0100-0200')."""
    return [int(c) if c.isdigit() else c.lower() for c in re.split(r'(\d+)', str(text))]


class DatasetVideo:

    def __init__(self, full_path: Path, root_dir: Path):
        self.full_path = full_path.resolve()
        self.root_dir = root_dir.resolve()
        try:
            rel = self.full_path.relative_to(self.root_dir)
            parts = rel.parts
        except ValueError:
            parts = (self.full_path.name,)

        # Hierarchy breakdown
        self.section = parts[0] if len(parts) > 1 else "Root"
        self.date = parts[1] if len(parts) > 2 else "General"
        self.session = parts[2] if len(parts) > 3 else "Session"
        self.filename = self.full_path.name
        self.rel_path_str = str(rel).replace("\\", "/")

        # Status & Results
        self.status = "Pending"  # Pending, In Progress, Completed, Failed
        self.vehicles_detected = 0
        self.mean_ctr_off = 0.0
        self.std_ctr_off = 0.0
        self.error_reason = ""
        self.timestamp = ""
        self.output_vid = ""
        self.output_csv = ""
        self.output_xlsx = ""

    def to_dict(self) -> dict:
        return {
            "rel_path": self.rel_path_str,
            "full_path": str(self.full_path),
            "section": self.section,
            "date": self.date,
            "session": self.session,
            "filename": self.filename,
            "status": self.status,
            "vehicles_detected": self.vehicles_detected,
            "mean_ctr_off": round(self.mean_ctr_off, 4),
            "std_ctr_off": round(self.std_ctr_off, 4),
            "error_reason": self.error_reason,
            "timestamp": self.timestamp,
            "output_vid": self.output_vid,
            "output_csv": self.output_csv,
            "output_xlsx": self.output_xlsx,
        }

    def from_dict(self, d: dict):
        self.status = d.get("status", "Pending")
        self.vehicles_detected = d.get("vehicles_detected", 0)
        self.mean_ctr_off = d.get("mean_ctr_off", 0.0)
        self.std_ctr_off = d.get("std_ctr_off", 0.0)
        self.error_reason = d.get("error_reason", "")
        self.timestamp = d.get("timestamp", "")
        self.output_vid = d.get("output_vid", "")
        self.output_csv = d.get("output_csv", "")
        self.output_xlsx = d.get("output_xlsx", "")


def dimensions_for_video(vid: DatasetVideo, road_width: float,
                         lane_width: float, rect_width: float,
                         rect_length: float) -> Tuple[float, float, float, float]:
    """Return the geometry used for this section/date calibration group."""
    if (vid.section, vid.date) in THREE_LANE_GROUPS:
        return 10.5, 3.5, 10.5, rect_length
    return road_width, lane_width, rect_width, rect_length


def calibration_matches(path: Path, rect_width: float,
                        rect_length: float) -> bool:
    """Reject an existing calibration made for different real-world dimensions."""
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return (abs(float(data.get("rw", -1)) - rect_width) < 1e-6 and
                abs(float(data.get("rl", -1)) - rect_length) < 1e-6)
    except (OSError, ValueError, TypeError, json.JSONDecodeError):
        return False


def discover_hierarchy(root_dir: Path) -> List[DatasetVideo]:
    """
    Traverse root_dir in strict hierarchy:
    Section -> Date -> Session -> Video
    Sorted naturally at each level.
    """
    if not root_dir.exists():
        sys.exit(f"[ERROR] Root directory does not exist: {root_dir}")

    discovered: List[DatasetVideo] = []

    # Get sections (LHS, RHS, etc.)
    sections = sorted([p for p in root_dir.iterdir() if p.is_dir()], key=lambda p: natural_key(p.name))
    if not sections:
        # Fallback if root itself contains video files directly
        vids = sorted([p for p in root_dir.glob("*") if p.suffix.lower() in VIDEO_EXTS], key=lambda p: natural_key(p.name))
        for v in vids:
            discovered.append(DatasetVideo(v, root_dir))
        return discovered

    for sec_path in sections:
        dates = sorted([p for p in sec_path.iterdir() if p.is_dir()], key=lambda p: natural_key(p.name))
        if not dates:
            # Check if videos exist directly in section folder
            for v in sorted([p for p in sec_path.glob("*") if p.suffix.lower() in VIDEO_EXTS], key=lambda p: natural_key(p.name)):
                discovered.append(DatasetVideo(v, root_dir))
            continue

        for date_path in dates:
            sessions = sorted([p for p in date_path.iterdir() if p.is_dir()], key=lambda p: natural_key(p.name))
            if not sessions:
                for v in sorted([p for p in date_path.glob("*") if p.suffix.lower() in VIDEO_EXTS], key=lambda p: natural_key(p.name)):
                    discovered.append(DatasetVideo(v, root_dir))
                continue

            for session_path in sessions:
                vids = sorted([p for p in session_path.iterdir() if p.suffix.lower() in VIDEO_EXTS], key=lambda p: natural_key(p.name))
                for v in vids:
                    discovered.append(DatasetVideo(v, root_dir))

    return discovered


class CheckpointManager:

    def __init__(self, output_dir: Path):
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.json_path = self.output_dir / "checkpoint.json"
        self.master_excel_path = self.output_dir / "master_dataset_report.xlsx"

    def load(self, discovered_videos: List[DatasetVideo]) -> List[DatasetVideo]:
        """Merge discovered videos with existing checkpoint."""
        checkpoint_data = {}
        if self.json_path.exists():
            try:
                with open(self.json_path, "r", encoding="utf-8") as f:
                    checkpoint_data = json.load(f).get("videos", {})
                print(f"[CHECKPOINT] Loaded existing state from {self.json_path}")
            except Exception as e:
                print(f"[WARN] Failed to load checkpoint: {e}")

        # A prior run may have a newer master report than the JSON checkpoint.
        # Restore completed rows from it before saving a new incremental report.
        if self.master_excel_path.exists() and XLSX_OK:
            try:
                wb = openpyxl.load_workbook(self.master_excel_path,
                                            read_only=True, data_only=True)
                ws = wb["Processing Manifest"]
                rows = ws.iter_rows(min_row=3, values_only=True)
                for row in rows:
                    rel_key = row[5]
                    status = row[6]
                    if not rel_key or status not in ("Completed", "Failed"):
                        continue
                    existing = checkpoint_data.get(rel_key, {})
                    if existing.get("status") in ("Completed", "Failed"):
                        continue
                    existing.update({
                        "status": status,
                        "vehicles_detected": row[7] or 0,
                        "mean_ctr_off": row[8] or 0.0,
                        "std_ctr_off": row[9] or 0.0,
                        "error_reason": row[10] or "",
                        "timestamp": row[11] or "",
                        "output_vid": row[13] or "",
                    })
                    output_dir = None
                    rel_parts = Path(rel_key).parts
                    if len(rel_parts) >= 4:
                        output_dir = self.output_dir.joinpath(*rel_parts[:3],
                                                             Path(rel_parts[-1]).stem)
                        stem = Path(rel_parts[-1]).stem
                        existing["output_csv"] = str(output_dir / f"{stem}_offsets.csv")
                        existing["output_xlsx"] = str(output_dir / f"{stem}_offsets.xlsx")
                    checkpoint_data[rel_key] = existing
                wb.close()
                print(f"[CHECKPOINT] Restored completed rows from {self.master_excel_path}")
            except Exception as e:
                print(f"[WARN] Failed to restore master report state: {e}")

        merged: List[DatasetVideo] = []
        for vid in discovered_videos:
            rel_key = vid.rel_path_str
            if rel_key in checkpoint_data:
                vid.from_dict(checkpoint_data[rel_key])
                # If program was interrupted while "In Progress", reset to "Pending" for safe retry
                if vid.status == "In Progress":
                    vid.status = "Pending"
                    vid.error_reason = "Interrupted during previous run — retrying"
            merged.append(vid)

        return merged

    def save(self, root_dir: Path, videos: List[DatasetVideo]):
        """Save JSON checkpoint and update Master Excel workbook."""
        data = {
            "root_dir": str(root_dir),
            "last_updated": datetime.datetime.now().isoformat(),
            "videos": {v.rel_path_str: v.to_dict() for v in videos}
        }
        with open(self.json_path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2)

        self._export_master_excel(videos)

    def _export_master_excel(self, videos: List[DatasetVideo]):
        if not XLSX_OK:
            return

        existing_raw_rows = []
        existing_raw_videos = set()
        if self.master_excel_path.exists():
            try:
                old_wb = openpyxl.load_workbook(self.master_excel_path,
                                                read_only=True, data_only=True)
                old_ws = old_wb["All Vehicles Raw Data"]
                old_rows = list(old_ws.iter_rows(values_only=True))
                if old_rows:
                    existing_raw_rows = old_rows[1:]
                    existing_raw_videos = {
                        row[4] for row in existing_raw_rows if len(row) > 4 and row[4]
                    }
                old_wb.close()
            except Exception as e:
                print(f"[WARN] Could not preserve existing raw data: {e}")

        wb = openpyxl.Workbook()
        thin = Side(style="thin", color="CCCCCC")
        bdr = Border(left=thin, right=thin, top=thin, bottom=thin)

        H_FILL = PatternFill("solid", start_color="0D1B40")  # Dark Header
        C_FILL = PatternFill("solid", start_color="E2F0D9")  # Green = Completed
        F_FILL = PatternFill("solid", start_color="FCE4D6")  # Red = Failed
        P_FILL = PatternFill("solid", start_color="FFF2CC")  # Yellow = In Progress
        N_FILL = PatternFill("solid", start_color="F2F2F2")  # Gray = Pending
        CENTER_ALIGN = Alignment(horizontal="center")
        LEFT_ALIGN = Alignment(horizontal="left")

        # ── Sheet 1: Processing Manifest ─────────────────────────────────
        ws = wb.active
        ws.title = "Processing Manifest"

        headers = [
            "#", "Section", "Date", "Time/Session", "Video Filename",
            "Relative Path", "Status", "Vehicles Detected",
            "Mean Center Off.(m)", "Std Center Off.(m)", "Error / Failure Reason",
            "Timestamp", "Full Path", "Annotated Video Path"
        ]
        col_widths = [5, 12, 12, 15, 25, 35, 14, 16, 18, 18, 30, 20, 45, 45]

        # Title Row
        ws.merge_cells(f"A1:{chr(64+len(headers))}1")
        tc = ws["A1"]
        tc.value = "Automated Video-Processing Master Manifest — BTP IIT Kharagpur"
        tc.font = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
        tc.fill = H_FILL
        tc.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        # Header Row
        ws.append(headers)
        ws.row_dimensions[2].height = 22
        for c_i, (h, cw) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=2, column=c_i)
            cell.font = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
            cell.fill = H_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = bdr
            ws.column_dimensions[openpyxl.utils.get_column_letter(c_i)].width = cw

        # Data Rows
        for i, v in enumerate(videos, 1):
            row_data = [
                i, v.section, v.date, v.session, v.filename,
                v.rel_path_str, v.status, v.vehicles_detected,
                v.mean_ctr_off if v.status == "Completed" else "",
                v.std_ctr_off if v.status == "Completed" else "",
                v.error_reason, v.timestamp, str(v.full_path), v.output_vid
            ]
            ws.append(row_data)

            # Fill based on status
            if v.status == "Completed":
                fill = C_FILL
            elif v.status == "Failed":
                fill = F_FILL
            elif v.status == "In Progress":
                fill = P_FILL
            else:
                fill = N_FILL

            for c_i in range(1, len(headers) + 1):
                cell = ws.cell(row=2+i, column=c_i)
                cell.font = Font(name="Calibri", size=10)
                cell.fill = fill
                cell.border = bdr
                cell.alignment = LEFT_ALIGN if c_i in (5, 6, 11, 13, 14) else CENTER_ALIGN

        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:{openpyxl.utils.get_column_letter(len(headers))}2"

        # ── Sheet 2: All Vehicles Raw Data ───────────────────────────────
        ws2 = wb.create_sheet("All Vehicles Raw Data")
        raw_headers = [
            "#", "Section", "Date", "Time/Session", "Video", "Frame",
            "Timestamp(s)", "Vehicle ID", "Vehicle Type",
            "Left Wheel Off(m)", "Right Wheel Off(m)", "Center Off(m)",
            "LW BEV-X(m)", "RW BEV-X(m)", "Confidence"
        ]
        ws2.append(raw_headers)
        for c_i, h in enumerate(raw_headers, 1):
            cell = ws2.cell(row=1, column=c_i)
            cell.font = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
            cell.fill = H_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = bdr

        raw_row_idx = 2
        for row_vals in existing_raw_rows:
            ws2.append(list(row_vals))
            for c_i in range(1, len(raw_headers) + 1):
                cell = ws2.cell(row=raw_row_idx, column=c_i)
                cell.font = Font(name="Calibri", size=9)
                cell.border = bdr
                cell.alignment = Alignment(horizontal="center")
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
            raw_row_idx += 1

        for v in videos:
            if (v.status != "Completed" or not v.output_csv or
                    not Path(v.output_csv).exists() or v.filename in existing_raw_videos):
                continue
            try:
                with open(v.output_csv, "r", encoding="utf-8") as csv_f:
                    reader = csv.DictReader(csv_f)
                    for r_i, r in enumerate(reader, 1):
                        row_vals = [
                            raw_row_idx - 1, v.section, v.date, v.session, v.filename,
                            r.get("Frame", ""), r.get("Timestamp(s)", ""),
                            r.get("VehicleID", ""), r.get("VehicleType", ""),
                            float(r.get("LeftWheelOffset(m)", 0.0)),
                            float(r.get("RightWheelOffset(m)", 0.0)),
                            float(r.get("CenterOffset(m)", 0.0)),
                            float(r.get("LeftWheelX_BEV(m)", 0.0)),
                            float(r.get("RightWheelX_BEV(m)", 0.0)),
                            float(r.get("Confidence", 0.0)),
                        ]
                        ws2.append(row_vals)
                        for c_i in range(1, len(raw_headers) + 1):
                            cell = ws2.cell(row=raw_row_idx, column=c_i)
                            cell.font = Font(name="Calibri", size=9)
                            cell.border = bdr
                            cell.alignment = Alignment(horizontal="center")
                            if isinstance(cell.value, float):
                                cell.number_format = "0.0000"
                        raw_row_idx += 1
            except Exception as e:
                print(f"[WARN] Failed to include CSV data for {v.filename}: {e}")

        for col in ws2.columns:
            ws2.column_dimensions[openpyxl.utils.get_column_letter(col[0].column)].width = 15
        ws2.freeze_panes = "A2"

        try:
            wb.save(self.master_excel_path)
        except Exception as e:
            print(f"[WARN] Could not write master Excel (file open?): {e}")


def process_video_single(vid: DatasetVideo, out_root: Path, default_calib_file: str,
                         model: str, conf: float, iou: float, imgsz: int,
                         skip_frames: int, no_riders: bool, road_width: float,
                         lane_width: float, rect_width: float, rect_length: float,
                         force_calib: bool = False, save_video: bool = False) -> Tuple[bool, str]:
    """Execute pipeline.py on a single video file with per-(Section, Date) calibration."""
    road_width, lane_width, rect_width, rect_length = dimensions_for_video(
        vid, road_width, lane_width, rect_width, rect_length)
    # Group key per Section and Date: e.g. LHS_1st_May, LHS_2nd_May, RHS_1st_May, RHS_2nd_May
    group_name = f"{vid.section}_{vid.date}".replace(" ", "_")
    group_calib = out_root / f"calib_{group_name}.json"
    legacy_calib = out_root / f"calib_{vid.section}.json"
    dimensions_changed = group_calib.exists() and not calibration_matches(
        group_calib, rect_width, rect_length)

    if group_calib.exists() and not force_calib and not dimensions_changed:
        calib_path = str(group_calib)
        is_headless = True
    elif vid.section == "LHS" and vid.date == "1st May" and legacy_calib.exists() and not force_calib:
        calib_path = str(legacy_calib)
        is_headless = True
    elif default_calib_file and Path(default_calib_file).exists() and not force_calib and (vid.section == "LHS" and vid.date == "1st May"):
        calib_path = default_calib_file
        is_headless = True
    else:
        # First video for this (Section, Date) group -> prompt user interactively
        calib_path = str(group_calib)
        is_headless = False
        print("\n" + "=" * 70, flush=True)
        print(f"  [INTERACTIVE CALIBRATION REQUIRED FOR: Section '{vid.section}' | Date '{vid.date}']", flush=True)
        print(f"  Video: {vid.rel_path_str}", flush=True)
        print("  Opening interactive calibration window...", flush=True)
        print("    Phase 1: Click 4 road corners (Far-Left, Far-Right, Near-Left, Near-Right)", flush=True)
        print("    Phase 2: Click 2 points on the RIGHT road boundary", flush=True)
        print("    Phase 3: Click 2 points across the road for the MEASUREMENT GATE", flush=True)
        print(f"  Saved to '{group_calib.name}' -> Reused headlessly for all '{vid.section} / {vid.date}' videos.", flush=True)
        print("=" * 70 + "\n", flush=True)

    # Output path structured by hierarchy: outputs/LHS/1st May/0000-0100/video_stem/
    out_dir = out_root / vid.section / vid.date / vid.session / vid.full_path.stem
    out_dir.mkdir(parents=True, exist_ok=True)
    out_vid_path = out_dir / f"{vid.full_path.stem}_annotated.mp4"

    cmd = [
        sys.executable, "pipeline.py",
        "--source",       str(vid.full_path),
        "--model",        model,
        "--output",       str(out_vid_path),
        "--conf",         str(conf),
        "--iou",          str(iou),
        "--imgsz",        str(imgsz),
        "--skip-frames",  str(skip_frames),
        "--road-width",   str(road_width),
        "--lane-width",   str(lane_width),
        "--rect-width",   str(rect_width),
        "--rect-length",  str(rect_length),
        "--calib-file",   calib_path,
    ]
    if is_headless:
        cmd.append("--headless")
    if not save_video:
        cmd.append("--no-video")
    if force_calib or dimensions_changed:
        cmd.append("--force-calib")
    if no_riders:
        cmd.append("--no-riders")

    vid.output_vid = str(out_vid_path) if save_video else "Not Saved"
    csv_file = out_dir / f"{vid.full_path.stem}_offsets.csv"
    xlsx_file = out_dir / f"{vid.full_path.stem}_offsets.xlsx"
    vid.output_csv = str(csv_file)
    vid.output_xlsx = str(xlsx_file)

    print(f"\n[{datetime.datetime.now().strftime('%H:%M:%S')}] RUNNING: {vid.rel_path_str}")
    proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace", bufsize=1)
    output_lines = []
    if proc.stdout is not None:
        for line in iter(proc.stdout.readline, ''):
            line_str = line.strip()
            if any(line_str.startswith(tag) for tag in ("[PROGRESS]", "[INFO]", "[CALIB]", "[DONE]", "[XLSX]", "[CSV]")):
                print(f"   └─ {line_str}", flush=True)
            output_lines.append(line_str)
    proc.wait()

    if proc.returncode != 0:
        err_msg = "\n".join(output_lines[-20:]) if output_lines else f"Process exited with code {proc.returncode}"
        return False, err_msg

    # Read offsets CSV to compute stats
    if csv_file.exists():
        try:
            with open(csv_file, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                rows = list(reader)
                vid.vehicles_detected = len(rows)
                if rows:
                    ctr_offs = [float(r["CenterOffset(m)"]) for r in rows if "CenterOffset(m)" in r]
                    if ctr_offs:
                        import numpy as np
                        vid.mean_ctr_off = float(np.mean(ctr_offs))
                        vid.std_ctr_off = float(np.std(ctr_offs))
        except Exception as e:
            print(f"[WARN] Error parsing output CSV: {e}")

    return True, ""


def print_summary(videos: List[DatasetVideo], current_index: int = -1):
    total = len(videos)
    completed = sum(1 for v in videos if v.status == "Completed")
    failed = sum(1 for v in videos if v.status == "Failed")
    pending = sum(1 for v in videos if v.status == "Pending")
    in_prog = sum(1 for v in videos if v.status == "In Progress")

    pct = (completed / total * 100.0) if total > 0 else 0.0

    print("\n" + "=" * 70)
    print("  AUTOMATED VIDEO-PROCESSING PIPELINE SUMMARY")
    print("=" * 70)
    print(f"  Total Videos Found   : {total}")
    print(f"  Completed            : {completed} ({pct:.1f}%)")
    print(f"  Failed               : {failed}")
    print(f"  Pending              : {pending}")
    if in_prog:
        print(f"  In Progress          : {in_prog}")
    if 0 <= current_index < total:
        print(f"  Current Position     : [{current_index + 1}/{total}] {videos[current_index].rel_path_str}")
    elif completed == total and total > 0:
        print("  Current Position     : ALL VIDEOS COMPLETED SUCCESSFULY!")
    print("=" * 70 + "\n")


def run_post_analysis(output_dir: Path):
    """Run analyze.py on all output CSVs."""
    csv_files = list(output_dir.glob("**/*_offsets.csv"))
    if not csv_files:
        print("[ANALYSIS] No CSV files found for analysis.")
        return

    analysis_dir = output_dir / "combined_analysis"
    cmd = [
        sys.executable, "analyze.py",
        "--csv"
    ] + [str(c) for c in csv_files] + [
        "--output-dir", str(analysis_dir)
    ]
    print(f"\n[ANALYSIS] Running combined statistical analysis on {len(csv_files)} CSV files...")
    res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
    print(res.stdout)
    print(f"[ANALYSIS] Analysis & plots saved to -> {analysis_dir.resolve()}")


def main():
    ap = argparse.ArgumentParser(
        description="Automated Resumable Video-Processing System — BTP IIT Kharagpur",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter
    )
    ap.add_argument("--root",          default=r"f:\BTP\NH-19 1st and 2nd May",
                    help="Root folder containing LHS / RHS hierarchy")
    ap.add_argument("--output-dir",   default="outputs",
                    help="Directory to save master outputs & checkpoint")
    ap.add_argument("--calib-file",   default="calib.json",
                    help="Calibration JSON file")
    ap.add_argument("--model",        default="yolo11n",
                    help="YOLO model name (yolo11n/yolo11s recommended for fast CPU execution, yolo11x for GPU)")
    ap.add_argument("--conf",         type=float, default=0.20,
                    help="Confidence threshold")
    ap.add_argument("--iou",          type=float, default=0.50,
                    help="NMS IoU threshold")
    ap.add_argument("--imgsz",        type=int, default=640,
                    help="Inference image resolution")
    ap.add_argument("--skip-frames",  type=int, default=2,
                    help="Process every Nth frame for speed")
    ap.add_argument("--no-riders",    action="store_true",
                    help="Exclude rider class 0")
    ap.add_argument("--road-width",   type=float, default=7.0)
    ap.add_argument("--lane-width",   type=float, default=3.5)
    ap.add_argument("--rect-width",   type=float, default=7.0)
    ap.add_argument("--rect-length",  type=float, default=8.0)
    ap.add_argument("--start-section", default="",
                    help="Start processing at this section/date group (for example RHS)")
    ap.add_argument("--start-date", default="",
                    help="Date paired with --start-section (for example '2nd May')")
    ap.add_argument("--reset-section", default="",
                    help="Reset this section/date group to Pending before processing")
    ap.add_argument("--reset-date", default="",
                    help="Date paired with --reset-section")
    ap.add_argument("--force-calib", action="store_true",
                    help="Redo calibration for the selected group")
    ap.add_argument("--save-video",   action="store_true",
                    help="Save annotated MP4 video files to disk (default False to save disk space and speed up processing)")
    ap.add_argument("--retry-failed", action="store_true",
                    help="Retry previously failed videos")
    args = ap.parse_args()

    root_dir = Path(args.root)
    out_root = Path(args.output_dir)

    # 1. Discover hierarchy
    discovered = discover_hierarchy(root_dir)
    print(f"[DISCOVERY] Found {len(discovered)} video(s) in strict hierarchy order.")

    # 2. Load Checkpoint
    chk_mgr = CheckpointManager(out_root)
    videos = chk_mgr.load(discovered)

    reset_group = (args.reset_section, args.reset_date)
    if any(reset_group):
        if not all(reset_group):
            sys.exit("[ERROR] --reset-section and --reset-date must be used together")
        reset_videos = [v for v in videos if (v.section, v.date) == reset_group]
        if not reset_videos:
            sys.exit(f"[ERROR] Reset group not found: {args.reset_section}/{args.reset_date}")
        for vid in reset_videos:
            vid.status = "Pending"
            vid.error_reason = ""
            vid.timestamp = ""
            vid.vehicles_detected = 0
            vid.mean_ctr_off = 0.0
            vid.std_ctr_off = 0.0
        chk_mgr.save(root_dir, videos)
        print(f"[RESET] Reset {len(reset_videos)} videos in {args.reset_section}/{args.reset_date}")

    # Signal handler for graceful exit on Ctrl+C
    def sig_handler(sig, frame):
        print("\n[INTERRUPT] Received signal to stop. Saving checkpoint & master report...")
        chk_mgr.save(root_dir, videos)
        print_summary(videos)
        sys.exit(0)

    signal.signal(signal.SIGINT, sig_handler)

    # Print initial summary
    print_summary(videos)

    start_group = (args.start_section, args.start_date)
    start_reached = not any(start_group)
    if any(start_group) and not any(
            (v.section, v.date) == start_group for v in videos):
        sys.exit(f"[ERROR] Start group not found: {args.start_section}/{args.start_date}")

    # 3. Sequential Processing Loop
    for idx, vid in enumerate(videos):
        if not start_reached:
            if (vid.section, vid.date) != start_group:
                continue
            start_reached = True
        if vid.status == "Completed":
            continue
        if vid.status == "Failed" and not args.retry_failed:
            continue

        # Mark In Progress & save checkpoint immediately
        vid.status = "In Progress"
        vid.timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        chk_mgr.save(root_dir, videos)

        # Run pipeline
        success, err = process_video_single(
            vid=vid,
            out_root=out_root,
            default_calib_file=args.calib_file,
            model=args.model,
            conf=args.conf,
            iou=args.iou,
            imgsz=args.imgsz,
            skip_frames=args.skip_frames,
            no_riders=args.no_riders,
            road_width=args.road_width,
            lane_width=args.lane_width,
            rect_width=args.rect_width,
            rect_length=args.rect_length,
            force_calib=args.force_calib,
            save_video=args.save_video,
        )

        vid.timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        if success:
            vid.status = "Completed"
            vid.error_reason = ""
            print(f"[SUCCESS] {vid.rel_path_str} -> Detected {vid.vehicles_detected} vehicles | Mean Ctr: {vid.mean_ctr_off:.3f}m")
        else:
            vid.status = "Failed"
            vid.error_reason = err
            print(f"[FAILED] {vid.rel_path_str} -> Error: {err}")

        # Save checkpoint & master Excel immediately after each video!
        chk_mgr.save(root_dir, videos)
        print_summary(videos, current_index=idx)

    # 4. Final Post-Processing & Figure Generation
    print("\n[COMPLETE] All videos processed!")
    chk_mgr.save(root_dir, videos)
    run_post_analysis(out_root)
    print_summary(videos)


if __name__ == "__main__":
    main()
